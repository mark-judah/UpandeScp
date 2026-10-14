"""Completed-plan detection, attribution reconciliation, report blocks and the Excel export."""

import io
import unittest
from datetime import datetime
from unittest import mock

import openpyxl

from upande_scp.serverscripts.reports import finance_api as API
from upande_scp.serverscripts.reports import finance_attribution as FAT
from upande_scp.serverscripts.reports import finance_paths as FP

KB = {"alias": {}, "kind": {"abamectin": "Insecticide"}, "ai_t": {"abamectin": {"Mites": 1}}, "prod_ai": {}, "prods": {}}
KINDS = {"Mites": "pest", "Thrips": "pest", "Mildew": "disease"}


def M(wo, day, wh="GH1", qty=1.2, item="MIX", h=8):
	return {"wo": wo, "name": f"M-{wo}", "dt": datetime(2026, 9, day, h), "item": item, "warehouse": wh, "qty": qty}


def I(name, day, wh="GH1", qty=1.2, item="MIX", h=12):
	return {"name": name, "dt": datetime(2026, 9, day, h), "item": item, "warehouse": wh, "qty": qty}


class TestCompletedPlans(unittest.TestCase):
	def test_manufacture_with_matching_issue_is_complete(self):
		self.assertEqual(FP.match_completed([M("A", 3)], [I("MI1", 3)])["A"]["name"], "MI1")

	def test_manufacture_without_issue_is_not_complete(self):
		self.assertEqual(FP.match_completed([M("A", 3)], []), {})

	def test_issue_from_another_greenhouse_or_quantity_does_not_match(self):
		self.assertEqual(FP.match_completed([M("A", 3)], [I("X", 3, wh="GH2")]), {})
		self.assertEqual(FP.match_completed([M("A", 3)], [I("X", 3, qty=1.0)]), {})

	def test_issue_before_the_manufacture_does_not_match(self):
		self.assertEqual(FP.match_completed([M("A", 5)], [I("X", 4)]), {})

	def test_issue_too_long_after_does_not_match(self):
		self.assertEqual(FP.match_completed([M("A", 1)], [I("X", 1)] and [{**I("X", 1), "dt": datetime(2026, 11, 1)}]), {})

	def test_each_issue_used_once_earliest_first(self):
		got = FP.match_completed([M("A", 3), M("B", 3, h=9)], [I("MI1", 3), I("MI2", 4)])
		self.assertEqual((got["A"]["name"], got["B"]["name"]), ("MI1", "MI2"))

	def test_second_plan_without_its_own_issue_stays_in_progress(self):
		got = FP.match_completed([M("A", 3), M("B", 3, h=9)], [I("MI1", 3)])
		self.assertEqual(set(got), {"A"})


def line(path="spray", code="P1", value=100.0, **kw):
	base = {"path": path, "date": "2026-09-10", "document": "D1", "issue": "I1", "work_order": "WO1" if path == "spray" else None,
	        "mr": None, "greenhouse": "GH1" if path == "spray" else None, "farm": "F", "cost_center": "CC1",
	        "item_code": code, "item_name": code, "qty": 1, "uom": "L", "value": value, "ledger": True,
	        "detail": "d", "plan_targets": ["Mites", "Mildew"] if path == "spray" else None, "recorded": []}
	base.update(kw)
	return base


PRODUCTS = {
	"P1": {"code": "P1", "name": "Chem", "category": "Chemical", "type": "Insecticide", "ingredients": ["abamectin"]},
	"P2": {"code": "P2", "name": "Plain chem", "category": "Chemical", "type": "Fungicide"},
	"F1": {"code": "F1", "name": "Foliar feed", "category": "Foliar"},
}
REFS = {"plans": {}, "mona_ai_t": {}}


def attr(lines):
	return FP.attribute(lines, REFS, KB, KINDS, PRODUCTS)


class TestAttribution(unittest.TestCase):
	def test_rows_sum_to_line_values_to_the_cent(self):
		lines = [line(code="P2", value=100.01), line(code="F1", value=33.33, plan_targets=["Mites", "Thrips", "Mildew"]),
		         line("direct", "P2", 50.0), line("direct", "F1", 20.0)]
		rows = attr(lines)
		self.assertAlmostEqual(sum(r["attributed"] for r in rows), sum(ln["value"] for ln in lines), places=6)

	def test_recorded_line_is_recorded_with_100_percent(self):
		rows = attr([line(code="P2", value=90, recorded=["Mildew"])])
		self.assertEqual([(r["target"], r["basis"], r["credibility"], r["accuracy"]) for r in rows],
		                 [("Mildew", "recorded", "recorded", 1.0)])

	def test_foliar_separate_and_equal_on_spray_farm_level_on_direct(self):
		rows = attr([line(code="F1", value=100), line("direct", "F1", 40)])
		spray = [r for r in rows if r["path"] == "spray"]
		self.assertEqual({(r["target"], r["attributed"]) for r in spray}, {("Mites", 50.0), ("Mildew", 50.0)})
		self.assertTrue(all(r["is_foliar"] for r in rows))
		d = [r for r in rows if r["path"] == "direct"]
		self.assertEqual([(r["target"], r["basis"]) for r in d], [(FP.FOLIAR_FARM, "foliar_farm")])

	def test_unattributed_direct_chemical(self):
		rows = attr([line("direct", "P2", 70)])
		self.assertEqual([(r["target"], r["basis"]) for r in rows], [(FP.UNATTRIBUTED, "unattributed")])


class TestBlocks(unittest.TestCase):
	def rows(self):
		return attr([
			line(code="P2", value=100, recorded=["Mites"]),
			line(code="P1", value=60),
			line(code="F1", value=40),
			line("direct", "P2", 25),
			line("direct", "F1", 15),
		])

	def test_foliar_never_in_chemical_totals(self):
		b = FP._block([r for r in self.rows() if r["path"] == "spray"])
		self.assertEqual(b["chemical"]["total"], 160.0)
		self.assertEqual(b["foliar"]["total"], 40.0)
		self.assertEqual(b["total"], 200.0)
		self.assertEqual(sum(p["value"] for p in b["chemical"]["by_pest"]), 160.0)

	def test_recorded_and_derived_are_separate_columns(self):
		b = FP._block([r for r in self.rows() if r["path"] == "spray"])
		mites = next(p for p in b["chemical"]["by_pest"] if p["target"] == "Mites")
		self.assertEqual(mites["recorded"], 100.0)
		self.assertEqual(mites["derived"], 60.0)  # P1 -> ingredient target Mites, band1_ai_target
		self.assertEqual(mites["value"], 160.0)
		self.assertAlmostEqual(b["chemical"]["derived_accuracy"], 0.89)

	def test_unattributed_separate(self):
		b = FP._block([r for r in self.rows() if r["path"] == "direct"])
		self.assertEqual(b["chemical"]["unattributed"], 25.0)
		self.assertEqual(b["foliar"]["farm_level"], 15.0)

	def test_filters(self):
		rows = self.rows()
		got, notes = FP.apply_filters(rows, paths=("direct",))
		self.assertTrue(all(r["path"] == "direct" for r in got))
		got, notes = FP.apply_filters(rows, paths=("spray", "direct"), greenhouse="GH1")
		self.assertTrue(all(r["path"] == "spray" for r in got) and notes)
		got, _ = FP.apply_filters(rows, pests=["Mites"])
		self.assertEqual({r["target"] for r in got}, {"Mites"})
		got, _ = FP.apply_filters(rows, pests=["Mites", FP.FOLIAR])
		self.assertTrue(any(r["is_foliar"] for r in got))


class TestExport(unittest.TestCase):
	def setUp(self):
		lines = [line(code="P2", value=100, recorded=["Mites"], detail="a"), line(code="P1", value=60, detail="b"),
		         line(code="F1", value=40, detail="c"), line("direct", "P2", 25, detail="d", mr="MR-1"),
		         line("direct", "F1", 15, detail="e")]
		self.data = {"rows": attr(lines), "lines": lines, "spray": {"in_progress": {"plans": 2, "value": 10.0}}, "kinds": KINDS}
		self.lines = lines

	def build(self, **kw):
		with mock.patch.object(FP, "collect_rows", return_value=self.data), \
		     mock.patch.object(FP, "ledger_total", side_effect=lambda ls: sum(x["value"] for x in ls)):
			content, sheets = API.build_workbook("2026-09-01", "2026-09-30", **kw)
		return openpyxl.load_workbook(io.BytesIO(content)), sheets

	def test_sheets_present_and_product_sheet_only_when_asked(self):
		wb, _ = self.build(per_product=False)
		self.assertEqual(wb.sheetnames, ["Summary", "By greenhouse", "Lines"])
		wb, _ = self.build(per_product=True)
		self.assertEqual(wb.sheetnames, ["Summary", "By greenhouse", "By product", "Lines"])

	def test_lines_sum_to_summary_totals(self):
		wb, _ = self.build()
		lines = wb["Lines"]
		head = [c.value for c in lines[1]]
		col = head.index("Attributed value")
		total_lines = round(sum(r[col] for r in lines.iter_rows(min_row=2, values_only=True)), 2)
		table = list(wb["Summary"].iter_rows(values_only=True))
		start = next(i for i, r in enumerate(table) if r[0] == "Totals by path")
		summ = {r[0]: r for r in table[start:start + 6] if r and r[0]}
		self.assertEqual(total_lines, round(sum(ln["value"] for ln in self.lines), 2))
		self.assertEqual(round(summ["Total"][7], 2), total_lines)
		self.assertEqual(summ["Spray plan"][7], 200.0)
		self.assertEqual(summ["Direct issue"][7], 40.0)

	def test_foliar_in_own_columns_and_credibility_columns(self):
		wb, _ = self.build()
		head = [c.value for c in wb["Lines"][1]]
		for h in ("Recorded value", "Derived value", "Foliar (equal split)", "Basis", "Credibility", "Est. accuracy"):
			self.assertIn(h, head)
		rows = list(wb["Lines"].iter_rows(min_row=2, values_only=True))
		bi, ci, ai = head.index("Basis"), head.index("Credibility"), head.index("Est. accuracy")
		labels = {r[ci] for r in rows}
		self.assertIn("Recorded", labels)
		self.assertTrue(any(x.startswith("Derived") and "est. 89%" in x for x in labels))
		self.assertIn("Foliar (equal split)", labels)
		rec = next(r for r in rows if r[bi] == "recorded")
		self.assertEqual(rec[ai], 1.0)

	def test_filters_echoed_and_path_filter(self):
		wb, _ = self.build(paths=("direct",), pests=["Mites"])
		text = [r for r in wb["Summary"].iter_rows(values_only=True)]
		flat = [c for r in text for c in r if c]
		self.assertIn("Direct issue", flat)
		self.assertIn("Mites", flat)
		self.assertNotIn("Spray plan", [r[0] for r in text if r[0] == "Spray plan"])

	def test_direct_is_farm_level_in_greenhouse_sheet(self):
		wb, _ = self.build()
		rows = list(wb["By greenhouse"].iter_rows(min_row=2, values_only=True))
		self.assertTrue(any(r[0] == "Direct issue" and r[2] == FP.FARM_LEVEL for r in rows))

	def test_header_rows_only_are_styled(self):
		wb, _ = self.build()
		ws = wb["Lines"]
		self.assertTrue(ws["A1"].font.bold)
		self.assertFalse(ws["A2"].font.bold)
		self.assertEqual(ws["M2"].number_format, "General")
